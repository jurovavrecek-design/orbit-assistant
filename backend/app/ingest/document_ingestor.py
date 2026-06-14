from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader


class DocumentIngestor:

    def load_documents(self):

        docs = []

        folders = [
            "docs/apex",
            "docs/userstories",
            "docs/pdf",
            "docs/docx",
            "docs/pptx"
        ]

        for folder in folders:

            if not Path(folder).exists():
                continue

            for file in Path(folder).glob("*"):

                if file.name.startswith(".") or file.name.startswith("~$"):
                    continue

                try:

                    #
                    # PDF
                    #
                    if file.suffix.lower() == ".pdf":

                        loader = PyPDFLoader(str(file))
                        pdf_docs = loader.load()

                        for doc in pdf_docs:

                            doc.metadata["source"] = file.name
                            doc.metadata["type"] = "pdf"

                        docs.extend(pdf_docs)

                        print(f"Loaded PDF {file.name}")

                        continue

                    #
                    # DOCX
                    #
                    if file.suffix.lower() == ".docx":

                        docx = DocxDocument(file)

                        text = "\n".join(
                            p.text
                            for p in docx.paragraphs
                        )

                        docs.append(
                            Document(
                                page_content=text,
                                metadata={
                                    "source": file.name,
                                    "type": "docx"
                                }
                            )
                        )

                        print(f"Loaded DOCX {file.name}")

                        continue

                    #
                    # PPTX
                    #
                    if file.suffix.lower() == ".pptx":

                        prs = Presentation(file)

                        for slide_number, slide in enumerate(prs.slides, start=1):

                            slide_title = ""

                            try:
                                if (
                                    len(slide.shapes) > 0
                                    and hasattr(slide.shapes[0], "text")
                                ):
                                    slide_title = slide.shapes[0].text.strip()

                            except:
                                pass

                            slide_text = f"SLIDE TITLE: {slide_title}\n\n"

                            for shape in slide.shapes:

                                if hasattr(shape, "text"):

                                    txt = shape.text.strip()

                                    if txt:
                                        slide_text += txt + "\n"

                            lower_slide = slide_text.lower()

                            skip_keywords = [
                                "module objectives",
                                "knowledge check",
                                "exercise"
                            ]

                            if any(keyword in lower_slide for keyword in skip_keywords):
                                continue

                            if slide_text.strip():

                                docs.append(
                                    Document(
                                        page_content=slide_text,
                                        metadata={
                                            "source": file.name,
                                            "type": "pptx",
                                            "slide": slide_number
                                        }
                                    )
                                )

                        print(f"Loaded PPTX {file.name}")

                        continue

                    text = file.read_text(errors="ignore")

                    docs.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": file.name,
                                "type": file.suffix.lower()
                            }
                        )
                    )

                    print(f"Loaded {file.name}")

                except Exception as e:

                    print(f"Failed {file.name}: {e}")

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=1500,
            chunk_overlap=300
        )

        final_docs = []

        for doc in docs:

            if doc.metadata.get("type") == "pptx":

                final_docs.append(doc)

            else:

                chunks = splitter.split_documents([doc])

                final_docs.extend(chunks)

        for i, chunk in enumerate(final_docs):

            chunk.metadata["chunk"] = i

        print(f"Created {len(final_docs)} chunks")

        return final_docs